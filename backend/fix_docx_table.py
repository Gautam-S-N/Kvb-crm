from docx import Document

doc = Document('D:/kvb-crm/backend/src/assets/dryer_template.docx')

changed = False

for t in doc.tables:
    for row in t.rows:
        # We know the specific row has 5 cells
        if len(row.cells) == 5:
            rate_cell = row.cells[3]
            if '{totalAmt}/-' in rate_cell.text:
                # Need to replace the `{totalAmt}` with `{unitPrice}`
                for p in rate_cell.paragraphs:
                    if '{totalAmt}' in p.text:
                        for run in p.runs:
                            if '{totalAmt}' in run.text:
                                run.text = run.text.replace('{totalAmt}', '{unitPrice}')
                                changed = True
                        if not changed:
                            p.text = p.text.replace('{totalAmt}', '{unitPrice}')
                            changed = True

            units_cell = row.cells[2]
            if 'Set' in units_cell.text:
                for p in units_cell.paragraphs:
                    if 'Set' in p.text:
                         for run in p.runs:
                             if 'Set' in run.text:
                                 run.text = run.text.replace('Set', '{units}')
                                 changed = True
                         if not changed:
                             p.text = p.text.replace('Set', '{units}')
                             changed = True
        
if changed:
    doc.save('D:/kvb-crm/backend/src/assets/dryer_template.docx')
    print("Fixed docx table formatting: replaced rate cell with {unitPrice} and units cell with {units}")
else:
    print("No changes made.")
