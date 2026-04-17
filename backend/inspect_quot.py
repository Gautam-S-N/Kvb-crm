from docx import Document
from docx.shared import Pt, Inches, Cm
from docx.oxml.ns import qn
import os

doc = Document('D:/kvb-crm/solar-tunnel-dryer-quotation.docx')

print("=== PARAGRAPHS ===")
for i, para in enumerate(doc.paragraphs):
    if para.text.strip():
        print(f"[{i}] style='{para.style.name}' align={para.alignment} text={repr(para.text[:120])}")

print("\n=== TABLES ===")
for ti, table in enumerate(doc.tables):
    print(f"\nTable {ti}: {len(table.rows)} rows x {len(table.columns)} cols")
    for ri, row in enumerate(table.rows):
        for ci, cell in enumerate(row.cells):
            txt = cell.text.strip()
            if txt:
                print(f"  [{ri},{ci}]: {repr(txt[:80])}")
