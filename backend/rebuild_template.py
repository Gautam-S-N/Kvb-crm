"""
FULL CLEAN REBUILD from scratch:
1. Start from ORIGINAL Dryer format for CRM.docx (never been touched by our scripts)
2. Apply text replacements using python-docx
3. Patch in New_Header&Footer1.docx header/footer using ElementTree
4. Output dryer_template.docx
"""
import zipfile, shutil, re
from xml.etree import ElementTree as ET
from docx import Document

ORIG_DRYER = 'D:/kvb-crm/Dryer format for CRM.docx'
BASE_HF    = 'D:/kvb-crm/New_Header&Footer1.docx'
OUT        = 'D:/kvb-crm/backend/src/assets/dryer_template.docx'
TEMP_DOCX  = 'D:/kvb-crm/backend/dryer_tagged.docx'
TEMP_ZIP   = OUT + '.tmp'

# ─── STEP 1: Apply text tags using python-docx ────────────────────────────────
replacements = {
    '13.04.2026': '{qtnDate}',
    'Mr. Chakaresh': '{toName}',
    'QTN.KVB.STD.005. A.080426 Solar Tunnel Dryer for 20w x 54L = 1080 Sq ft': '{subjectLine}',
    'Rectangular type with top parabolic Shape': '{productType}',
    '54ft L X 20 ft W X 8.5 ft H': '{dimensions}',
    '8.5.5 feet': '{centerHeight}',
    'GP Square Pipe Frame 25x25mm': '{structureDoor}',
    'GP Square Pipe 40mm x 40mm': '{purlin}',
    'GP Square pipe 40x40mm': '{arch}',
    'Tray size 2ftx3ft \u2013 Customer Scope': '{traySize}',
    'Supply and installation of Polycarbonate sheet covered Solar Tunnel Dryer 1080 Sq ft.': '{itemDesc}',
    '5,83,200': '{totalAmt}',
    '70% Advance along with PO 30% against Performa invoice after inspection at factory prior to despatch': '{paymentTerms}',
    'To your account': '{deliveryTerms}',
    'Packing \u2013 3% extra (Bubble sheet / corrugated sheet)': '{packingTerms}',
    'Freight and insurance \u2013 To your account': '{freightTerms}',
    'presently 18%': '{gstRate}',
    'Five Lakh Eighty Three Thousand Two Hundred Rupees Only': '{amountWords}',
    'Five Lakhs Eight Three Thousand Two Hundred Rupees Only': '{amountWords}',
}

def replace_in_para(para):
    full = ''.join(r.text for r in para.runs)
    changed = False
    for old, new in replacements.items():
        if old in full:
            full = full.replace(old, new)
            changed = True
    if changed:
        for i, r in enumerate(para.runs):
            # Remove yellow highlight
            from docx.oxml.ns import qn
            rpr = r._r.find(qn('w:rPr'))
            if rpr is not None:
                for tag in ['w:highlight', 'w:shd']:
                    el = rpr.find(qn(tag))
                    if el is not None:
                        rpr.remove(el)
            r.text = full if i == 0 else ''

def strip_highlight(para):
    from docx.oxml.ns import qn
    for r in para.runs:
        rpr = r._r.find(qn('w:rPr'))
        if rpr is not None:
            for tag in [qn('w:highlight'), qn('w:shd')]:
                el = rpr.find(tag)
                if el is not None:
                    rpr.remove(el)

doc = Document(ORIG_DRYER)

for p in doc.paragraphs:
    replace_in_para(p)
    strip_highlight(p)

for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                replace_in_para(p)
                strip_highlight(p)

# Fix qty=01 -> {qty} (it's in a table cell)
for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                full = ''.join(r.text for r in p.runs)
                if full.strip() == '01':
                    for i, r in enumerate(p.runs):
                        r.text = '{qty}' if i == 0 else ''

doc.save(TEMP_DOCX)
print('Step 1 done: tagged docx saved')

# ─── STEP 2: Read the tagged docx at zip level ───────────────────────────────
def read_zip(path):
    with zipfile.ZipFile(path, 'r') as z:
        return {n: z.read(n) for n in z.namelist()}

tagged_files = read_zip(TEMP_DOCX)
base_files   = read_zip(BASE_HF)

tagged_doc_xml = tagged_files['word/document.xml'].decode('utf-8')
tagged_rels_xml = tagged_files['word/_rels/document.xml.rels'].decode('utf-8')

# ─── STEP 3: Extract body content (stop before final sectPr) ─────────────────
body_match = re.search(r'<w:body>(.*?)(<w:sectPr[\s>])', tagged_doc_xml, re.DOTALL)
body_content = body_match.group(1)

# Double-check: strip any remaining highlights at XML level too
body_content = re.sub(r'<w:highlight[^/]*/>', '', body_content)
body_content = re.sub(r'<w:shd[^/]*/>', '', body_content)

# ─── STEP 4: Figure out image rIds in tagged body ────────────────────────────
# python-docx preserves original rIds from ORIG_DRYER
blips = list(set(re.findall(r'r:embed="([^"]+)"', body_content)))
print('Image rIds embedded in body:', blips)

# Map what each one points to in tagged_rels
NS_R = 'http://schemas.openxmlformats.org/package/2006/relationships'
tagged_rels_root = ET.fromstring(tagged_rels_xml)
body_img_map = {}  # rId -> original media filename
for rel in tagged_rels_root:
    if rel.attrib.get('Id') in blips:
        target = rel.attrib['Target']  # e.g. media/image1.jpg
        body_img_map[rel.attrib['Id']] = target
        print(f"  {rel.attrib['Id']} -> {target}")

# ─── STEP 5: Remap body image rIds to avoid collision with base doc rIds ─────
# Base doc uses rId9->header1, rId10->header2, rId11->footer, rId12->header3
# We must rename body images to safe new IDs
rid_remap = {}
for i, rid in enumerate(blips):
    new_rid = f'rIdBodyImg{i+1}'
    rid_remap[rid] = new_rid
    body_content = body_content.replace(f'r:embed="{rid}"', f'r:embed="{new_rid}"')
    print(f'  Remapped {rid} -> {new_rid}')

# ─── STEP 6: Build new document.xml (tagged open + body + base sectPr) ────────
tagged_open = re.match(r'(.*?)<w:body>', tagged_doc_xml, re.DOTALL).group(0)

base_doc_xml = base_files['word/document.xml'].decode('utf-8')
sect_match = re.search(r'(<w:sectPr[\s>].*?</w:sectPr>)', base_doc_xml, re.DOTALL)
base_sectPr = sect_match.group(1)

new_doc = tagged_open + body_content + base_sectPr + '</w:body></w:document>'

# ─── STEP 6: Build new rels: all base rels + body image rels + hyperlinks ─────
base_rels_root = ET.fromstring(base_files['word/_rels/document.xml.rels'])
all_rels = []
for rel in base_rels_root:
    all_rels.append(dict(rel.attrib))

# Add body image rels (using remapped IDs)
for rel in tagged_rels_root:
    rid = rel.attrib.get('Id', '')
    if rid in blips:
        new_attrs = dict(rel.attrib)
        new_attrs['Id'] = rid_remap[rid]
        all_rels.append(new_attrs)

# Add hyperlinks from tagged doc
for rel in tagged_rels_root:
    if 'hyperlink' in rel.attrib.get('Type', ''):
        all_rels.append(dict(rel.attrib))

def make_rels_xml(rel_list):
    entries = []
    for a in rel_list:
        parts = ' '.join(f'{k}="{v}"' for k, v in a.items())
        entries.append(f'<Relationship {parts}/>')
    return ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\r\n'
            f'<Relationships xmlns="{NS_R}">'
            + ''.join(entries) + '</Relationships>')

new_rels = make_rels_xml(all_rels)

# ─── STEP 7: Content-Types ───────────────────────────────────────────────────
new_ct = base_files['[Content_Types].xml'].decode('utf-8')
for ext, ct in [('jpg', 'image/jpeg'), ('jpeg', 'image/jpeg'), ('png', 'image/png')]:
    if f'Extension="{ext}"' not in new_ct:
        new_ct = new_ct.replace('</Types>', f'<Default Extension="{ext}" ContentType="{ct}"/></Types>')

# ─── STEP 8: Assemble zip ─────────────────────────────────────────────────────
# Build a safe rename map: dryer body images get stored as body_img1.jpg etc.
# This GUARANTEES no collision with base media (image1.png, image2.jpeg, image3.png)
body_img_rename = {}  # original target path -> new safe path
body_img_data   = {}  # new safe path -> bytes

for i, (old_rid, old_target) in enumerate(body_img_map.items()):
    # old_target = e.g. "media/image1.jpg"
    ext = old_target.rsplit('.', 1)[-1]
    new_target = f'media/body_img{i+1}.{ext}'
    body_img_rename[old_target] = new_target
    # Also update the rels entry we already built to point to new filename
    for rel_attrs in all_rels:
        if rel_attrs.get('Id') == rid_remap[old_rid]:
            rel_attrs['Target'] = new_target
    # Store the actual image bytes
    src_key = f'word/{old_target}'
    if src_key in tagged_files:
        body_img_data[f'word/{new_target}'] = tagged_files[src_key]
    print(f'  Body image: {old_target} -> {new_target}')

# Rebuild rels XML with updated targets
new_rels = make_rels_xml(all_rels)

overrides = {
    'word/document.xml':            new_doc.encode('utf-8'),
    'word/_rels/document.xml.rels': new_rels.encode('utf-8'),
    '[Content_Types].xml':          new_ct.encode('utf-8'),
    'word/styles.xml':              tagged_files.get('word/styles.xml', base_files.get('word/styles.xml', b'')),
    'word/numbering.xml':           tagged_files.get('word/numbering.xml', b''),
}
# Add renamed body images (NEVER overwrite base media)
overrides.update(body_img_data)

with zipfile.ZipFile(TEMP_ZIP, 'w', zipfile.ZIP_DEFLATED) as zout:
    # Write ALL base files (header/footer/media intact)
    for name, data in base_files.items():
        if name not in overrides:
            zout.writestr(name, data)
    # Write overrides/additions
    for name, data in overrides.items():
        if data:
            zout.writestr(name, data)


shutil.move(TEMP_ZIP, OUT)
print('\nFinal template written to:', OUT)

# ─── Verify ──────────────────────────────────────────────────────────────────
with zipfile.ZipFile(OUT) as z:
    rels_check = z.read('word/_rels/document.xml.rels').decode()
    doc_check  = z.read('word/document.xml').decode()
    
    print('\nVerification:')
    print('  header1 in rels:', 'header1.xml' in rels_check)
    print('  footer1 in rels:', 'footer1.xml' in rels_check)
    
    for bid in blips:
        present = bid in rels_check
        print(f'  body img {bid} in rels: {present}')
    
    text = re.sub(r'<[^>]+>', ' ', doc_check)
    leftovers = re.findall(r'\{[a-z][a-zA-Z]+\}', text)
    print('  Unfilled tags:', leftovers)
    print('  Highlights:', 'w:highlight' in doc_check)
    print('  Media in zip:', [f for f in z.namelist() if 'media' in f])
