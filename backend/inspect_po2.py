import openpyxl
from openpyxl.utils import get_column_letter

wb = openpyxl.load_workbook('D:/kvb-crm/PO Format.xlsx', data_only=True)
ws = wb.active

print("Merged cells:", ws.merged_cells)
print()
print("Column widths:")
for col in ws.column_dimensions:
    print(f"  Col {col}: width={ws.column_dimensions[col].width}")
print()
print("Row heights:")
for row in ws.row_dimensions:
    print(f"  Row {row}: height={ws.row_dimensions[row].height}")
print()
print("All cells with values:")
for row in ws.iter_rows():
    for cell in row:
        if cell.value is not None:
            print(f"  {cell.coordinate} (col {cell.column}, row {cell.row}): {repr(str(cell.value)[:50])}")
