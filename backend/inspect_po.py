import openpyxl
from openpyxl.styles import PatternFill

wb = openpyxl.load_workbook('D:/kvb-crm/PO Format.xlsx', data_only=True)
ws = wb.active
print('Sheet:', ws.title, '  Dims:', ws.dimensions)
print()

for row in ws.iter_rows():
    for cell in row:
        val = cell.value
        fill = cell.fill
        rgb = 'none'
        if fill and fill.patternType and fill.patternType != 'none':
            try:
                rgb = fill.fgColor.rgb
            except:
                rgb = 'err'
        if val is not None:
            yellow = '*YELLOW*' if rgb in ('FFFFFF00','FFFF00','00FFFF00') else ''
            print(f'  {cell.coordinate}: {str(val)[:70]:70s} fill={rgb} {yellow}')
