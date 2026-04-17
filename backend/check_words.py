from docx import Document
import re

doc = Document('D:/kvb-crm/backend/src/assets/dryer_template.docx')

found = False
for p in doc.paragraphs:
    if 'Five' in p.text and 'Lakh' in p.text:
        print(f"Found in paragraph: '{p.text}'")
        found = True
        
for t in doc.tables:
    for row in t.rows:
        for cell in row.cells:
            if 'Five' in cell.text and 'Lakh' in cell.text:
                print(f"Found in table cell: '{cell.text}'")
                found = True

if not found:
    print("Not found in paragraphs or tables!")
