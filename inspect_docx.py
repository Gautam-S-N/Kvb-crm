import docx

def dump_all_tables(filepath):
    doc = docx.Document(filepath)
    for ti, table in enumerate(doc.tables):
        print(f"\n=== TABLE {ti} ===")
        for ri, row in enumerate(table.rows):
            row_texts = []
            for ci, cell in enumerate(row.cells):
                text = ' '.join(p.text for p in cell.paragraphs).strip()
                row_texts.append(f"C{ci}:[{text}]")
            print(f"  R{ri}: " + " | ".join(row_texts))

    print("\n=== PARAGRAPHS (non-empty) ===")
    for i, para in enumerate(doc.paragraphs):
        if para.text.strip():
            print(f"  P{i}: {para.text[:120]}")

if __name__ == '__main__':
    dump_all_tables('SOLAR PARABOLIC COOKER.docx')
